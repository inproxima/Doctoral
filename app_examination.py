import os
import re
import shutil
from datetime import datetime
from io import BytesIO

import streamlit as st
from openai import OpenAI
from docx import Document
from dotenv import load_dotenv
import PyPDF2


load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

OPENAI_MODEL = "gpt-5.5"
REASONING_EFFORT = "high"

st.set_page_config(
    page_title="Doctoral Dissertation Review",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
    /* Tighten tab label spacing */
    .stTabs [data-baseweb="tab-list"] { gap: 4px; }
    .stTabs [data-baseweb="tab"] {
        padding: 0.55rem 0.9rem;
        font-weight: 500;
    }
    /* Slightly larger section headers */
    h1, h2, h3 { letter-spacing: -0.01em; }
    /* Monospace for download filename hints */
    .stDownloadButton button { font-family: var(--font-monospace); }
    /* Calmer separator */
    hr { margin: 0.75rem 0 1rem 0; opacity: 0.4; }
    </style>
    """,
    unsafe_allow_html=True,
)


CHAPTER_DESCRIPTIONS = {
    1: ("Introduction", "📖", "first chapter of the dissertation (Introduction)"),
    2: ("Literature Review", "📚", "second chapter of the dissertation (Literature Review)"),
    3: ("Methodology", "🔬", "third chapter of the dissertation (Methodology)"),
    4: ("Results/Findings", "📊", "fourth chapter of the dissertation (Results/Findings)"),
    5: ("Discussion/Analysis", "💬", "fifth chapter of the dissertation (Discussion/Analysis)"),
    6: ("Conclusion", "🎯", "sixth chapter of the dissertation (Conclusion)"),
}

# Each chapter is judged only against its own remit. Stating the exclusions explicitly is what
# stops the same criterion — and therefore the same criticism — being generated six times.
CHAPTER_SCOPE = {
    1: {
        "owns": (
            "Framing of the research problem; statement and justification of the knowledge gap; "
            "research questions/objectives and their alignment with the problem; significance and "
            "rationale; scope, delimitations, and definitions of key terms; the roadmap of the argument."
        ),
        "not_here": (
            "Exhaustive synthesis or critique of the literature (Chapter 2); detailed methodological "
            "design, instruments, or sampling (Chapter 3); any findings (Chapter 4)."
        ),
    },
    2: {
        "owns": (
            "Coverage, currency, and selection logic of the literature; critical synthesis rather than "
            "summary; construction and justification of the theoretical/conceptual framework; how the "
            "reviewed work substantiates the gap claimed in Chapter 1; treatment of contradictory evidence."
        ),
        "not_here": (
            "Re-stating the problem or research questions (Chapter 1); operational detail of instruments, "
            "sampling, or analysis procedure (Chapter 3); interpretation of this study's results (Chapter 5)."
        ),
    },
    3: {
        "owns": (
            "Research design and its fit to the research questions; philosophical/paradigmatic positioning; "
            "population, sampling strategy, and sample size justification; instrumentation and its "
            "validation; data-collection procedure; analysis plan and its appropriateness; validity, "
            "reliability/trustworthiness, and threats to them; ethics, consent, and data governance; "
            "replicability of the described procedure."
        ),
        "not_here": (
            "Justifying the gap or the research questions themselves (Chapters 1–2); reporting or "
            "interpreting results (Chapters 4–5)."
        ),
    },
    4: {
        "owns": (
            "Completeness and accuracy of results reporting; alignment of reported results to each research "
            "question; correct and sufficient statistical/analytical reporting (effect sizes, assumptions, "
            "uncertainty, qualitative evidence and its provenance); quality, labelling, and self-sufficiency "
            "of tables and figures; restraint from interpretation; transparency about missing or negative results."
        ),
        "not_here": (
            "Re-explaining or re-justifying the method (Chapter 3); interpreting findings against the "
            "literature or drawing implications (Chapter 5)."
        ),
    },
    5: {
        "owns": (
            "Interpretation of findings against the framework and literature from Chapter 2; whether claims "
            "are proportionate to the evidence reported in Chapter 4; engagement with alternative or "
            "disconfirming explanations; articulation of the theoretical and practical contribution; "
            "honest, specific limitations and their consequences for the claims."
        ),
        "not_here": (
            "Re-reporting the results themselves (Chapter 4); re-describing the method (Chapter 3); "
            "final conclusions and future-research agenda (Chapter 6)."
        ),
    },
    6: {
        "owns": (
            "Whether the conclusions are warranted by, and traceable to, the reported findings; the explicit "
            "statement of the contribution to knowledge and its defensibility as doctoral-level; the "
            "closing of the loop back to the Chapter 1 problem and research questions; recommendations for "
            "practice/policy and their grounding; the future-research agenda; overall closure of the argument."
        ),
        "not_here": (
            "Introducing new data, analysis, or literature; repeating the interpretive argument of "
            "Chapter 5; restating limitations already given in Chapter 5 beyond a one-line reference."
        ),
    },
}

SEVERITY_SCALE = """Severity labels (use these exact words):
- **Critical** — undermines the defensibility of the doctoral claim; the degree cannot be awarded until it is resolved.
- **Major** — a substantive flaw in rigour, evidence, or argument that requires significant rework.
- **Moderate** — a real weakness that a demanding examiner would require addressed, but that does not threaten the core claim.
- **Minor** — a small gap in clarity, completeness, or convention."""

COVERAGE_TAG = "COVERAGE"
LEDGER_TAG = "LEDGER"


def _llm(prompt: str, system: str = "You are a senior doctoral examiner at a research-intensive university.") -> str:
    try:
        resp = client.chat.completions.create(
            model=OPENAI_MODEL,
            reasoning_effort=REASONING_EFFORT,
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": prompt},
            ],
        )
        return resp.choices[0].message.content
    except Exception as e:
        return f"Error calling {OPENAI_MODEL}: {e}"


def _split_block(text: str, tag: str) -> tuple[str, str]:
    """Split a delimited machine-readable block off the end of an LLM response.

    Returns (visible_text, block_body). A missing or malformed block yields
    (text, "") so a formatting slip degrades the carry-forward rather than the app.
    """
    if not text:
        return "", ""
    match = re.search(
        rf"<<<\s*{tag}\s*>>>(.*?)(?:<<<\s*END(?:\s*{tag})?\s*>>>|\Z)",
        text,
        re.DOTALL | re.IGNORECASE,
    )
    if not match:
        return text, ""
    body = "\n".join(
        line for line in match.group(1).splitlines() if not line.strip().startswith("```")
    ).strip()
    visible = (text[: match.start()] + text[match.end() :]).strip()
    return visible, body


def _parse_ledger(body: str) -> list[dict]:
    """Best-effort parse of `ID | theme | severity | statement` rows. Display only."""
    rows = []
    for line in body.splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) < 4 or not re.fullmatch(r"C\d+-\d+", cells[0], re.IGNORECASE):
            continue
        rows.append(
            {
                "ID": cells[0].upper(),
                "Theme": cells[1],
                "Severity": cells[2],
                "Issue": " | ".join(cells[3:]),
            }
        )
    return rows


def _prior_coverage(chapter_num: int) -> str:
    parts = [
        f"Chapter {n} ({CHAPTER_DESCRIPTIONS[n][0]}) criteria already cover:\n"
        + st.session_state[f"ch{n}_coverage"].strip()
        for n in range(1, chapter_num)
        if st.session_state.get(f"ch{n}_coverage", "").strip()
    ]
    return "\n\n".join(parts)


def _prior_ledger(chapter_num: int) -> str:
    parts = [
        f"From Chapter {n} ({CHAPTER_DESCRIPTIONS[n][0]}):\n"
        + st.session_state[f"ch{n}_ledger"].strip()
        for n in range(1, chapter_num)
        if st.session_state.get(f"ch{n}_ledger", "").strip()
    ]
    return "\n\n".join(parts)


def _full_ledger() -> str:
    return _prior_ledger(7)


def _unreviewed_before(chapter_num: int) -> list[int]:
    return [n for n in range(1, chapter_num) if f"ch{n}_review" not in st.session_state]


def _recurrence_map() -> dict[str, list[int]]:
    """Chapters that referenced an issue raised in a different chapter."""
    hits: dict[str, list[int]] = {}
    for n in range(1, 7):
        review = st.session_state.get(f"ch{n}_review", "")
        for origin, seq in set(re.findall(r"\bC(\d+)-(\d+)\b", review)):
            if int(origin) != n:
                hits.setdefault(f"C{origin}-{seq}", []).append(n)
    return {k: sorted(v) for k, v in hits.items()}


def _normalize_text(text: str) -> str:
    text = re.sub(r"(\w+)-\n(\w+)", r"\1\2", text)
    text = re.sub(r"(?<!\n\s)\n(?!\s\n)", " ", text.strip())
    text = re.sub(r"\n\s*\n", "\n\n", text)
    return text


def save_uploaded_files(uploaded_files):
    temp_dir = "TempDir"
    os.makedirs(temp_dir, exist_ok=True)
    for uploaded_file in uploaded_files:
        if uploaded_file is not None:
            file_path = os.path.join(temp_dir, uploaded_file.name)
            with open(file_path, "wb") as f:
                f.write(uploaded_file.getbuffer())


def convert_pdf_to_text(directory_path: str = "TempDir") -> dict:
    directory_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "TempDir")
    pdf_texts = {}
    for filename in os.listdir(directory_path):
        if filename.endswith(".pdf"):
            filepath = os.path.join(directory_path, filename)
            with open(filepath, "rb") as file:
                reader = PyPDF2.PdfReader(file)
                pages = [p.extract_text() for p in reader.pages if p.extract_text()]
                pdf_texts[filename] = _normalize_text("\n".join(pages))
    return pdf_texts


def convert_docx_to_text(directory_path: str = "TempDir") -> dict:
    directory_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "TempDir")
    docx_texts = {}
    for filename in os.listdir(directory_path):
        if filename.endswith(".docx"):
            filepath = os.path.join(directory_path, filename)
            document = Document(filepath)
            text = "\n".join(p.text for p in document.paragraphs)
            docx_texts[filename] = _normalize_text(text)
    return docx_texts


def extract_text_from_uploaded_file(uploaded_file) -> str:
    if uploaded_file is None:
        return ""
    filename = uploaded_file.name.lower()
    try:
        if filename.endswith(".pdf"):
            reader = PyPDF2.PdfReader(BytesIO(uploaded_file.getvalue()))
            pages = [p.extract_text() for p in reader.pages if p.extract_text()]
            combined = "\n".join(pages)
        elif filename.endswith(".docx"):
            document = Document(BytesIO(uploaded_file.getvalue()))
            combined = "\n".join(p.text for p in document.paragraphs)
        elif filename.endswith(".txt"):
            combined = uploaded_file.getvalue().decode("utf-8", errors="ignore")
        else:
            return ""
        return _normalize_text(combined)
    except Exception:
        return ""


def clear_chapters_folder():
    chapters_dir = "chapters"
    if os.path.exists(chapters_dir):
        for filename in os.listdir(chapters_dir):
            file_path = os.path.join(chapters_dir, filename)
            try:
                if os.path.isfile(file_path):
                    os.unlink(file_path)
                elif os.path.isdir(file_path):
                    shutil.rmtree(file_path)
            except Exception as e:
                st.error(f"Error deleting {file_path}: {e}")
        st.success("Chapters folder cleared.")
    else:
        st.warning("Chapters folder does not exist.")


def generate_chapter_criteria(
    text: str, dissertation_title: str, chapter_num: int, prior_coverage: str
) -> str:
    _, _, descriptor = CHAPTER_DESCRIPTIONS[chapter_num]
    scope = CHAPTER_SCOPE[chapter_num]
    already_covered = (
        f"""
CRITERIA ALREADY ESTABLISHED IN EARLIER CHAPTERS — do NOT restate or rephrase any of these.
The examination applies them once, where they belong.
{prior_coverage}
"""
        if prior_coverage
        else ""
    )
    return _llm(
        f"""You are an examiner of a doctoral dissertation titled "{dissertation_title}".
This is the {descriptor}:

{text}

THIS CHAPTER'S REMIT — every criterion you generate must fall inside it:
{scope["owns"]}

OUT OF SCOPE — these belong to other chapters and will be judged there. Generate NO criteria for them:
{scope["not_here"]}
{already_covered}
Generate rigorous evaluation criteria you will use to judge this chapter, and only this chapter.
Each criterion must be one this chapter alone can satisfy or fail — if a criterion could be applied
unchanged to any other chapter of any dissertation, it is too generic: drop it or sharpen it until
it is specific to the remit above.

For each criterion give:
- a short title (5 words or fewer)
- a one-sentence description
- an explicit benchmark for what a doctoral-quality version of THIS chapter demonstrates

After the criteria, output this block exactly, listing each criterion title on its own line with no
numbering, bullets, or commentary:

<<<{COVERAGE_TAG}>>>
Criterion title
Criterion title
<<<END>>>"""
    )


def generate_chapter_review(
    text: str, criteria: str, dissertation_title: str, chapter_num: int, prior_ledger: str
) -> str:
    _, _, descriptor = CHAPTER_DESCRIPTIONS[chapter_num]
    scope = CHAPTER_SCOPE[chapter_num]
    if prior_ledger:
        recurrence_rules = f"""
ISSUES ALREADY RAISED IN EARLIER CHAPTERS OF THIS EXAMINATION:
{prior_ledger}

De-duplication rules — these are binding:
- If one of the issues above also manifests in this chapter, do NOT argue it again. List it under
  "Recurring Issues" as its ID plus ONE line of new, chapter-specific evidence. Nothing more.
- "Points to Consider" must contain only issues that are NEW to this chapter and not already
  covered by an ID above, even if worded differently. An issue that is merely a rephrasing of an
  existing ID is a recurrence, not a new point.
- If this chapter shows an existing issue is in fact resolved here, say so in one line under
  "Recurring Issues".
"""
        recurring_section = (
            "## 🔁 Recurring Issues\n"
            "References only, one line each, in the form `[C2-03] — new evidence in this chapter: …`. "
            "No restatement of the issue, no rationale, no remedy. Write "
            "`None — no earlier issue recurs in this chapter.` if that is the case."
        )
    else:
        recurrence_rules = ""
        recurring_section = (
            "## 🔁 Recurring Issues\n"
            "Write `Not applicable — this is the first chapter reviewed.`"
        )

    return _llm(
        f"""You are an examiner of a doctoral dissertation titled "{dissertation_title}".
This is the {descriptor}:

{text}

THIS CHAPTER'S REMIT — judge it on this and nothing else:
{scope["owns"]}

OUT OF SCOPE — belongs to other chapters and is judged there. Do not raise it as a fault of this chapter:
{scope["not_here"]}

Evaluation criteria to apply:
{criteria}
{recurrence_rules}
{SEVERITY_SCALE}

Produce a critical, evidence-based review using EXACTLY these four Markdown headings, in this order:

## ✅ Strengths
Concrete strengths, each tied to a specific section, passage, table, or argument. No generic praise.

## ⚠️ Points to Consider
Only issues new to this chapter, RANKED most important first (Critical, then Major, then Moderate,
then Minor). Number them in rank order. Format each as:

### C{chapter_num}-01 · [Severity] Short title
- **Where:** section / page / table reference
- **Evidence:** what the text actually does or fails to do
- **Why it matters:** the consequence for the defensibility of the doctoral claim
- **Resolution:** what a satisfactory version would look like

Use sequential IDs C{chapter_num}-01, C{chapter_num}-02, … in the ranked order.

{recurring_section}

## ✏️ Edit Points
A single Markdown table of every concrete change required in this chapter — substantive rewrites and
surface corrections alike — in this format:

| Location | Required edit | Ref |
|---|---|---|
| §{chapter_num}.2, p.47 | Rewrite the passage to state X explicitly and justify it | C{chapter_num}-02 |
| §{chapter_num}.4, Table 2 | Add a caption and state the units | — |

Ref cites the Point to Consider the edit discharges, or `—` for a standalone surface fix
(typo, grammar, citation style, figure/table labelling, inconsistent terminology).

Each fact appears exactly ONCE in this review: a Point to Consider carries the diagnosis and
rationale, its Edit Point carries only the concrete instruction. Do not repeat the rationale in the
table, and do not repeat the edit wording in the Points to Consider.

Finally, output this block exactly, one row per Point to Consider raised above (not recurrences),
pipe-delimited, no header row and no table separator line:

<<<{LEDGER_TAG}>>>
C{chapter_num}-01 | theme in 2-4 words | Severity | one-line statement of the issue
<<<END>>>

Maintain a fair but demanding standard appropriate for a research-intensive doctoral examination."""
    )


def generate_chapter_oral_questions(text: str, dissertation_title: str, chapter_num: int) -> str:
    _, _, descriptor = CHAPTER_DESCRIPTIONS[chapter_num]
    return _llm(
        f"""You are an examiner in a doctoral oral defence for the dissertation titled "{dissertation_title}".
This is the {descriptor}:

{text}

Generate a numbered list of critical oral-defence questions. Probe the candidate's:
- Conceptual command of the material
- Methodological choices and trade-offs
- Awareness of alternative interpretations
- Limitations and threats to validity
- Connection of this chapter to the broader dissertation argument

Where useful, include short follow-up probes after a primary question."""
    )


def generate_overall_oral_questions(chapter_texts: dict, dissertation_title: str) -> str:
    bundle = "\n\n".join(
        f"Chapter {n} ({CHAPTER_DESCRIPTIONS[n][0]}):\n{chapter_texts[n]}"
        for n in range(1, 7)
        if chapter_texts.get(n)
    )
    return _llm(
        f"""You are an examiner in a doctoral oral defence for the dissertation titled "{dissertation_title}".

Full dissertation content by chapter:
{bundle}

Generate a numbered list of overarching oral-defence questions that assess:
- Coherence and progression of argument across chapters
- Significance and originality of the contribution to knowledge
- The candidate's grasp of broader implications and positioning
- Ethical considerations and study limitations
- Future research directions and the candidate's readiness for independent scholarship

Include short follow-up probes where useful."""
    )


def generate_dissertation_summary(chapter_texts: dict, dissertation_title: str) -> str:
    bundle = "\n\n".join(
        f"Chapter {n} ({CHAPTER_DESCRIPTIONS[n][0]}):\n{chapter_texts[n]}"
        for n in range(1, 7)
        if chapter_texts.get(n)
    )
    return _llm(
        f"""You are a doctoral examiner reviewing the dissertation titled "{dissertation_title}".

Full dissertation content by chapter:
{bundle}

Produce a comprehensive, well-structured summary covering:
1. Research Problem & Objectives
2. Research Gap & Significance
3. Theoretical Framework
4. Methodology Overview
5. Key Findings
6. Theoretical Contributions
7. Practical Implications
8. Limitations
9. Future Research
10. Overall Assessment of coherence, rigor, and contribution to knowledge

Be precise and objective."""
    )


def generate_final_review(
    chapter_reviews: dict, dissertation_title: str, chair_concerns: str, ledger: str
) -> str:
    blocks = "\n\n".join(
        f"Chapter {n} Review ({CHAPTER_DESCRIPTIONS[n][0]}):\n{chapter_reviews[n]}"
        for n in range(1, 7)
        if chapter_reviews.get(n)
    )
    focus = chair_concerns.strip() or "(no specific focus provided)"
    ledger_block = ledger.strip() or "(no issue ledger available)"
    return _llm(
        f"""You are an examiner producing the final assessment of the doctoral dissertation titled "{dissertation_title}".

Examination focus / chair concerns: {focus}

Consolidated issue ledger from every chapter (ID | theme | severity | statement):
{ledger_block}

Chapter-level reviews (source evidence):
{blocks}

Your task is to MERGE these into one assessment, not to concatenate them. Binding rules:
- Cluster issues that share a root cause into a SINGLE dissertation-level point, even when they were
  raised in different chapters under different wording. Cite every ID you have folded into the point.
- An issue that recurred across chapters is stated ONCE, listing the chapters it affects.
- Nothing in the final list may be a restatement of something already listed above it.
- Every point must trace to evidence in the chapter reviews. Do not invent new issues.

{SEVERITY_SCALE}

Rank the Points to Consider by importance using this rubric, applied in order:
1. Threat to the defensibility of the doctoral claim (severity).
2. Breadth — how many chapters the issue affects, and whether it propagates downstream.
3. Effort to remedy, used only to break ties among otherwise equal points.

Use EXACTLY these four Markdown headings, in this order:

## ✅ Strengths of the Dissertation
The genuine strengths of the work as a whole, each tied to specific chapters and evidence. Merge
strengths that were noted in several chapters into one statement.

## ⚠️ Points to Consider (ranked)
A single numbered list ordered most important first across the WHOLE dissertation — not grouped by
chapter. Format each as:

### 1. [Severity] Short title
- **Chapters affected:** e.g. 2, 5
- **Folded from:** the issue IDs merged into this point
- **Evidence:** the specific basis in the dissertation
- **Why it matters:** consequence for the doctoral claim
- **Required action:** what the candidate must do to resolve it

## ✏️ Edit Points
Every concrete change required, consolidated and grouped under a `### Chapter N` subheading, as a
Markdown table `| Location | Required edit | Ref |`. Ref cites the number of the Point to Consider
above that the edit discharges, or `—` for a standalone surface fix. Do not repeat rationale here —
the instruction only. Drop duplicate edits that appeared in more than one chapter review.

## 🏁 Overall Judgement
One of: acceptable as-is / minor revisions / major revisions / not acceptable for doctoral conferral.
Give the rationale in a short paragraph, referring to the ranked points by number.

Be evidence-based and direct."""
    )


def generate_decision_letter(dissertation_title: str, final_review: str, chair_concerns: str) -> str:
    focus = chair_concerns.strip() or "(no specific focus provided)"
    return _llm(
        f"""You are the chair of a doctoral examination committee. Draft a formal decision letter to the candidate for the dissertation titled "{dissertation_title}".

Examination focus / chair concerns: {focus}

Final assessment to base the letter on:
{final_review}

The letter must include:
1. Examination Decision (accept / minor revisions / major revisions / not acceptable)
2. Decision Rationale (concise, grounded in the assessment)
3. Summary of examiner concerns (organized by theme)
4. Required Actions (specific, numbered, with priority) if revisions are required
5. Recommendations for strengthening the work
6. Timeline for revisions if applicable

Draw the Required Actions directly from the assessment's ranked "Points to Consider", preserving
that rank order. Do not re-derive, re-rank, or expand the list, and do not repeat an action that a
higher-numbered action already covers.

Start with "Dear Candidate," and end with "Sincerely, [Committee Chair] on behalf of the Examination Committee."

Maintain a professional, constructive, and direct tone."""
    )


def generate_revised_dissertation_assessment(
    revised_text: str,
    dissertation_title: str,
    criteria_bundle: str,
    prior_final_review: str,
    chair_concerns: str,
) -> str:
    focus = chair_concerns.strip() or "(no specific focus provided)"
    return _llm(
        f"""You are conducting a second-round examination of the revised doctoral dissertation titled "{dissertation_title}".

Examination focus / chair concerns: {focus}

Prior evaluation criteria (all chapters):
{criteria_bundle}

Prior final assessment:
{prior_final_review}

Revised dissertation (full text):
{revised_text}

Produce a structured second-round report:
1. Resolved Items — issues now satisfactorily addressed, with specific evidence from the revised text.
2. Partially Resolved Items — what improved and what still falls short.
3. Unresolved or New Issues — original concerns still unmet plus any new concerns introduced by the revisions.
4. Overall Judgement — concise recommendation (accept / minor revisions / major revisions / not acceptable).

Be precise and cite sections where possible."""
    )


def evaluate_candidate_response(
    response_text: str,
    prior_final_review: str,
    revised_text: str,
    criteria_bundle: str,
    chair_concerns: str,
) -> str:
    focus = chair_concerns.strip() or "(no specific focus provided)"
    return _llm(
        f"""You are reviewing the candidate's point-by-point response letter for a doctoral dissertation revision.

Examination focus / chair concerns: {focus}

Prior evaluation criteria (all chapters):
{criteria_bundle}

Prior final assessment (issues to be addressed):
{prior_final_review}

Candidate's response letter:
{response_text}

Revised dissertation (evidence):
{revised_text}

Produce a structured adequacy matrix. For each substantive issue from the prior assessment:
- Issue (concise restatement)
- Candidate's Response Summary
- Evidence in the Revised Dissertation (section / page where possible)
- Adequacy Rating: Fully Addressed / Partially Addressed / Not Addressed
- Notes

Conclude with an Overall Adequacy Score and a list of remaining changes still required."""
    )


with st.sidebar:
    st.header("🎓 Examination Setup")
    dissertation_title = st.text_input(
        "Dissertation Title",
        help="Used throughout the examination process.",
    )
    chair_concerns = st.text_area(
        "Examination Focus / Chair Concerns",
        placeholder="e.g. methodological rigor, theoretical contribution, statistical robustness…",
        help="Optional. Concerns flagged by the committee chair that should shape the review.",
    )
    if dissertation_title:
        st.caption(f"**Current dissertation:** *{dissertation_title}*")
    st.markdown("---")

    with st.expander("📁 File Management", expanded=True):
        uploaded_files = st.file_uploader(
            "Upload chapter files (PDF or DOCX)",
            type=["pdf", "docx"],
            accept_multiple_files=True,
            help="Name files clearly (e.g. 'chapter_1.pdf') for automatic assignment.",
        )

        col_a, col_b = st.columns(2)
        with col_a:
            process_clicked = st.button("Process Files", type="primary", width="stretch")
        with col_b:
            if st.button("Clear Chapters", width="stretch"):
                clear_chapters_folder()

        if process_clicked:
            if uploaded_files:
                if os.path.exists("TempDir"):
                    shutil.rmtree("TempDir")
                save_uploaded_files(uploaded_files)
                pdf_texts = convert_pdf_to_text()
                docx_texts = convert_docx_to_text()
                all_texts = {**pdf_texts, **docx_texts}
                st.session_state.extracted_texts = all_texts

                # New source text invalidates every derived artifact — a stale ledger would
                # otherwise be carried forward into the new examination.
                for n in range(1, 7):
                    for suffix in ("criteria", "coverage", "review", "ledger", "oral_questions"):
                        st.session_state.pop(f"ch{n}_{suffix}", None)
                for key in (
                    "dissertation_summary",
                    "final_review",
                    "final_decision_letter",
                    "overall_oral_questions",
                    "revised_assessment",
                    "candidate_response_evaluation",
                    "review_date",
                ):
                    st.session_state.pop(key, None)

                for filename in all_texts.keys():
                    lowercase_name = filename.lower()
                    for n in range(1, 7):
                        if f"chapter_{n}" in lowercase_name or f"chapter{n}" in lowercase_name:
                            st.session_state[f"chapter{n}_file"] = filename

                st.success(f"{len(uploaded_files)} file(s) processed.")
                assigned = [
                    f"Chapter {n}: {st.session_state[f'chapter{n}_file']}"
                    for n in range(1, 7)
                    if f"chapter{n}_file" in st.session_state
                ]
                if assigned:
                    st.caption("Auto-assigned chapters:")
                    for line in assigned:
                        st.caption(f"• {line}")
            else:
                st.error("Please upload files first.")

    with st.expander("ℹ️ About"):
        st.caption(
            f"Powered by **OpenAI {OPENAI_MODEL}** with `reasoning_effort=\"{REASONING_EFFORT}\"`."
        )
        st.caption(
            "AI-assisted doctoral examination. Final decisions must always involve human expert judgement."
        )


st.title("🎓 Doctoral Dissertation Examination")
if dissertation_title:
    st.caption(f"**Dissertation:** {dissertation_title}  |  **Focus:** {chair_concerns.strip() or '*not specified*'}")
else:
    st.caption("Enter the dissertation title in the sidebar to begin.")
st.markdown("---")


tab_labels = [
    f"{icon} Chapter {n} ({name})" for n, (name, icon, _) in CHAPTER_DESCRIPTIONS.items()
] + [
    "📝 Final Review & Decision",
    "🗣️ Oral Defence",
    "🔁 Revision Review",
    "📈 Analytics",
]
tabs = st.tabs(tab_labels)
chapter_tabs = tabs[:6]
final_tab, oral_tab, revision_tab, analytics_tab = tabs[6], tabs[7], tabs[8], tabs[9]


def render_chapter_tab(tab, chapter_num: int):
    name, icon, _ = CHAPTER_DESCRIPTIONS[chapter_num]
    with tab:
        st.header(f"{icon} Chapter {chapter_num} — {name}")

        if "extracted_texts" not in st.session_state or not st.session_state.extracted_texts:
            st.info("Upload and process chapter files from the sidebar to begin.")
            return

        file_options = list(st.session_state.extracted_texts.keys())
        default_index = 0
        chapter_file_key = f"chapter{chapter_num}_file"
        if chapter_file_key in st.session_state:
            try:
                default_index = file_options.index(st.session_state[chapter_file_key])
            except ValueError:
                default_index = 0

        selected_file = st.selectbox(
            f"Select file for Chapter {chapter_num}",
            file_options,
            index=default_index,
            key=f"select_ch{chapter_num}",
        )

        if not selected_file:
            return

        with st.expander("📄 Extracted text preview"):
            preview = st.session_state.extracted_texts[selected_file]
            preview = preview[:3000] + "…" if len(preview) > 3000 else preview
            st.text_area(
                "Preview",
                preview,
                height=260,
                key=f"text_ch{chapter_num}",
                label_visibility="collapsed",
            )

        prior_coverage = _prior_coverage(chapter_num)
        prior_ledger = _prior_ledger(chapter_num)

        if chapter_num > 1:
            missing = _unreviewed_before(chapter_num)
            if missing:
                st.warning(
                    "Chapter"
                    + ("s " if len(missing) > 1 else " ")
                    + ", ".join(str(m) for m in missing)
                    + " not yet reviewed — de-duplication context is partial. "
                    "Review chapters in order for best results."
                )
            with st.expander("📎 Carried forward from earlier chapters"):
                if not prior_coverage and not prior_ledger:
                    st.caption("Nothing carried forward yet.")
                if prior_coverage:
                    st.markdown("**Criteria already covered** — not regenerated for this chapter")
                    st.code(prior_coverage, language=None)
                if prior_ledger:
                    st.markdown("**Issues already raised** — referenced, not restated, in this chapter")
                    st.code(prior_ledger, language=None)

        has_criteria = f"ch{chapter_num}_criteria" in st.session_state
        col1, col2 = st.columns(2)
        with col1:
            if st.button(
                "Generate Criteria",
                type="primary" if not has_criteria else "secondary",
                key=f"btn_criteria_ch{chapter_num}",
                width="stretch",
            ):
                if not dissertation_title:
                    st.error("Enter the dissertation title in the sidebar first.")
                else:
                    with st.status(f"Generating criteria for Chapter {chapter_num}…", expanded=True) as status:
                        criteria = generate_chapter_criteria(
                            st.session_state.extracted_texts[selected_file],
                            dissertation_title,
                            chapter_num,
                            prior_coverage,
                        )
                        visible, coverage = _split_block(criteria, COVERAGE_TAG)
                        st.session_state[f"ch{chapter_num}_criteria"] = visible
                        st.session_state[f"ch{chapter_num}_coverage"] = coverage
                        status.update(label="Criteria ready.", state="complete")
                    has_criteria = True

        with col2:
            if st.button(
                "Generate Review",
                type="primary" if has_criteria else "secondary",
                key=f"btn_review_ch{chapter_num}",
                width="stretch",
                disabled=not has_criteria,
            ):
                with st.status(f"Generating review for Chapter {chapter_num}…", expanded=True) as status:
                    review = generate_chapter_review(
                        st.session_state.extracted_texts[selected_file],
                        st.session_state[f"ch{chapter_num}_criteria"],
                        dissertation_title,
                        chapter_num,
                        prior_ledger,
                    )
                    visible, ledger = _split_block(review, LEDGER_TAG)
                    st.session_state[f"ch{chapter_num}_review"] = visible
                    st.session_state[f"ch{chapter_num}_ledger"] = ledger
                    status.update(label="Review ready.", state="complete")

        if has_criteria:
            with st.expander("📋 Evaluation Criteria", expanded=True):
                st.markdown(st.session_state[f"ch{chapter_num}_criteria"])
            st.download_button(
                "Download Criteria",
                data=st.session_state[f"ch{chapter_num}_criteria"],
                file_name=f"chapter_{chapter_num}_criteria.txt",
                mime="text/plain",
                key=f"dl_criteria_ch{chapter_num}",
            )

        if f"ch{chapter_num}_review" in st.session_state:
            with st.expander("🔍 Chapter Review", expanded=True):
                st.markdown(st.session_state[f"ch{chapter_num}_review"])
            st.download_button(
                "Download Review",
                data=st.session_state[f"ch{chapter_num}_review"],
                file_name=f"chapter_{chapter_num}_review.txt",
                mime="text/plain",
                key=f"dl_review_ch{chapter_num}",
            )


for n, tab in zip(range(1, 7), chapter_tabs):
    render_chapter_tab(tab, n)


with final_tab:
    st.header("📝 Final Review & Decision")
    st.caption("Synthesizes the dissertation summary, comprehensive final review, and a formal decision letter.")

    if "extracted_texts" not in st.session_state or not st.session_state.extracted_texts:
        st.info("Upload and process chapter files first.")
    else:
        file_options = list(st.session_state.extracted_texts.keys())
        st.subheader("Chapter assignments for synthesis")
        summary_files = {}
        cols = st.columns(3)
        for n in range(1, 7):
            col = cols[(n - 1) % 3]
            with col:
                default_index = 0
                if f"chapter{n}_file" in st.session_state:
                    try:
                        default_index = file_options.index(st.session_state[f"chapter{n}_file"])
                    except ValueError:
                        default_index = 0
                summary_files[n] = st.selectbox(
                    f"Chapter {n} file",
                    file_options,
                    index=default_index,
                    key=f"final_select_ch{n}",
                )

        st.markdown("---")
        st.subheader("Generate artifacts")
        col_s, col_f, col_d = st.columns(3)
        chapter_texts = {n: st.session_state.extracted_texts.get(summary_files[n], "") for n in range(1, 7)}

        with col_s:
            if st.button("Dissertation Summary", type="primary", width="stretch"):
                if not dissertation_title:
                    st.error("Enter the dissertation title in the sidebar first.")
                elif not all(chapter_texts[n] for n in range(1, 7)):
                    st.error("All six chapter files must be available.")
                else:
                    with st.status("Generating dissertation summary…", expanded=True) as status:
                        st.session_state["dissertation_summary"] = generate_dissertation_summary(
                            chapter_texts, dissertation_title
                        )
                        status.update(label="Summary ready.", state="complete")

        with col_f:
            chapter_reviews_ready = all(f"ch{n}_review" in st.session_state for n in range(1, 7))
            if st.button(
                "Final Comprehensive Review",
                type="primary" if chapter_reviews_ready else "secondary",
                width="stretch",
                disabled=not chapter_reviews_ready,
                help=None if chapter_reviews_ready else "Generate reviews for Chapters 1–6 first.",
            ):
                with st.status("Synthesizing final review…", expanded=True) as status:
                    chapter_reviews = {
                        n: st.session_state.get(f"ch{n}_review", "Not reviewed.") for n in range(1, 7)
                    }
                    st.session_state["final_review"] = generate_final_review(
                        chapter_reviews, dissertation_title, chair_concerns, _full_ledger()
                    )
                    st.session_state["review_date"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    status.update(label="Final review ready.", state="complete")

        with col_d:
            decision_ready = "final_review" in st.session_state
            if st.button(
                "Decision Letter",
                type="primary" if decision_ready else "secondary",
                width="stretch",
                disabled=not decision_ready,
                help=None if decision_ready else "Generate the final review first.",
            ):
                with st.status("Drafting decision letter…", expanded=True) as status:
                    st.session_state["final_decision_letter"] = generate_decision_letter(
                        dissertation_title,
                        st.session_state["final_review"],
                        chair_concerns,
                    )
                    status.update(label="Decision letter ready.", state="complete")

        if "dissertation_summary" in st.session_state:
            with st.expander("📑 Dissertation Summary", expanded=True):
                st.markdown(st.session_state["dissertation_summary"])
            st.download_button(
                "Download Summary",
                data=st.session_state["dissertation_summary"],
                file_name=f"{(dissertation_title or 'dissertation').replace(' ', '_')}_summary.txt",
                mime="text/plain",
                key="dl_summary",
            )

        if "final_review" in st.session_state:
            with st.expander("🧭 Final Comprehensive Review", expanded=True):
                st.markdown(st.session_state["final_review"])
            st.download_button(
                "Download Final Review",
                data=st.session_state["final_review"],
                file_name=f"{(dissertation_title or 'dissertation').replace(' ', '_')}_final_review.txt",
                mime="text/plain",
                key="dl_final_review",
            )

        if "final_decision_letter" in st.session_state:
            with st.expander("✉️ Decision Letter", expanded=True):
                st.markdown(st.session_state["final_decision_letter"])
            st.download_button(
                "Download Decision Letter",
                data=st.session_state["final_decision_letter"],
                file_name=f"{(dissertation_title or 'dissertation').replace(' ', '_')}_decision_letter.txt",
                mime="text/plain",
                key="dl_decision_letter",
            )

        if "final_decision_letter" in st.session_state:
            st.markdown("---")
            package = "\n\n".join(
                [
                    "DOCTORAL EXAMINATION PACKAGE",
                    f"Generated: {st.session_state.get('review_date', 'N/A')}",
                    f"Dissertation: {dissertation_title}",
                    f"Examination Focus: {chair_concerns.strip() or 'Not specified'}",
                    "=" * 50,
                    "DISSERTATION SUMMARY",
                    "=" * 50,
                    st.session_state.get("dissertation_summary", "N/A"),
                    "=" * 50,
                    "FINAL COMPREHENSIVE REVIEW",
                    "=" * 50,
                    st.session_state["final_review"],
                    "=" * 50,
                    "DECISION LETTER",
                    "=" * 50,
                    st.session_state["final_decision_letter"],
                ]
            )
            st.download_button(
                "📦 Download Complete Examination Package",
                data=package,
                file_name=f"{(dissertation_title or 'dissertation').replace(' ', '_')}_examination_package.txt",
                mime="text/plain",
                key="dl_package",
            )


with oral_tab:
    st.header("🗣️ Oral Defence Questions")
    st.caption("Per-chapter probing questions plus overarching defence questions.")

    if "extracted_texts" not in st.session_state or not st.session_state.extracted_texts:
        st.info("Upload and process chapter files first.")
    else:
        file_options = list(st.session_state.extracted_texts.keys())
        chapter_file_choices = {}

        for row_start in (1, 4):
            row_cols = st.columns(3)
            for offset, col in enumerate(row_cols):
                n = row_start + offset
                name, icon, _ = CHAPTER_DESCRIPTIONS[n]
                with col:
                    st.subheader(f"{icon} Chapter {n}")
                    default_index = 0
                    if f"chapter{n}_file" in st.session_state:
                        try:
                            default_index = file_options.index(st.session_state[f"chapter{n}_file"])
                        except ValueError:
                            default_index = 0
                    chapter_file_choices[n] = st.selectbox(
                        f"Chapter {n} file",
                        file_options,
                        index=default_index,
                        key=f"oral_select_ch{n}",
                        label_visibility="collapsed",
                    )
                    if st.button(
                        f"Generate Questions",
                        key=f"btn_oral_ch{n}",
                        width="stretch",
                    ):
                        if not dissertation_title:
                            st.error("Set dissertation title in sidebar.")
                        else:
                            with st.status(f"Generating Chapter {n} oral questions…", expanded=True) as status:
                                st.session_state[f"ch{n}_oral_questions"] = generate_chapter_oral_questions(
                                    st.session_state.extracted_texts[chapter_file_choices[n]],
                                    dissertation_title,
                                    n,
                                )
                                status.update(label="Questions ready.", state="complete")

                    if f"ch{n}_oral_questions" in st.session_state:
                        with st.expander(f"Chapter {n} questions"):
                            st.markdown(st.session_state[f"ch{n}_oral_questions"])
                        st.download_button(
                            "Download",
                            data=st.session_state[f"ch{n}_oral_questions"],
                            file_name=f"chapter_{n}_oral_questions.txt",
                            mime="text/plain",
                            key=f"dl_oral_ch{n}",
                            width="stretch",
                        )

        st.markdown("---")
        st.subheader("Overarching Defence Questions")
        all_files_set = all(chapter_file_choices.get(n) for n in range(1, 7))
        if st.button(
            "Generate Overarching Questions",
            type="primary",
            disabled=not all_files_set,
            help=None if all_files_set else "All six chapter files must be selected.",
        ):
            chapter_text_bundle = {
                n: st.session_state.extracted_texts[chapter_file_choices[n]] for n in range(1, 7)
            }
            with st.status("Generating overarching questions…", expanded=True) as status:
                st.session_state["overall_oral_questions"] = generate_overall_oral_questions(
                    chapter_text_bundle, dissertation_title
                )
                status.update(label="Questions ready.", state="complete")

        if "overall_oral_questions" in st.session_state:
            with st.expander("Overall dissertation questions", expanded=True):
                st.markdown(st.session_state["overall_oral_questions"])
            st.download_button(
                "Download Overarching Questions",
                data=st.session_state["overall_oral_questions"],
                file_name=f"{(dissertation_title or 'dissertation').replace(' ', '_')}_overall_oral_questions.txt",
                mime="text/plain",
                key="dl_overall_oral",
            )


with revision_tab:
    st.header("🔁 Revision Review")
    st.caption(
        "Upload the revised dissertation and the candidate's response letter to assess revision adequacy."
    )

    col_r1, col_r2 = st.columns(2)
    with col_r1:
        revised_file = st.file_uploader(
            "Revised dissertation (PDF or DOCX)",
            type=["pdf", "docx"],
            key="revised_dissertation_uploader",
            accept_multiple_files=False,
        )
    with col_r2:
        response_file = st.file_uploader(
            "Candidate response letter (PDF, DOCX, or TXT)",
            type=["pdf", "docx", "txt"],
            key="candidate_response_uploader",
            accept_multiple_files=False,
        )

    if st.button("Process Revision Files", type="primary"):
        revised_text = extract_text_from_uploaded_file(revised_file)
        response_text = extract_text_from_uploaded_file(response_file)
        if revised_text:
            st.session_state["revised_dissertation_text"] = revised_text
            st.success("Revised dissertation processed.")
        else:
            st.warning("No revised dissertation processed.")
        if response_text:
            st.session_state["candidate_response_text"] = response_text
            st.success("Candidate response processed.")
        else:
            st.info("Proceeding without a candidate response letter limits the adequacy evaluation.")

    if "revised_dissertation_text" in st.session_state:
        with st.expander("📄 Revised dissertation preview"):
            preview = st.session_state["revised_dissertation_text"]
            preview = preview[:3000] + "…" if len(preview) > 3000 else preview
            st.text_area("Preview", preview, height=260, key="revised_preview", label_visibility="collapsed")

    if "candidate_response_text" in st.session_state:
        with st.expander("✉️ Candidate response preview"):
            preview = st.session_state["candidate_response_text"]
            preview = preview[:3000] + "…" if len(preview) > 3000 else preview
            st.text_area("Preview", preview, height=260, key="response_preview", label_visibility="collapsed")

    st.markdown("---")
    st.subheader("Assess the Revision")

    criteria_bundle_parts = [
        f"Chapter {n} ({CHAPTER_DESCRIPTIONS[n][0]}) criteria:\n{st.session_state[f'ch{n}_criteria']}"
        for n in range(1, 7)
        if f"ch{n}_criteria" in st.session_state
    ]
    criteria_bundle = "\n\n".join(criteria_bundle_parts) or "(no prior criteria available)"
    prior_final_review = st.session_state.get("final_review", "(no prior final review available)")

    col_b1, col_b2 = st.columns(2)
    with col_b1:
        if st.button("Assess Revised Dissertation", type="primary", width="stretch"):
            if "revised_dissertation_text" not in st.session_state:
                st.error("Upload and process the revised dissertation first.")
            elif not criteria_bundle_parts:
                st.error("Generate per-chapter criteria in the chapter tabs first.")
            else:
                with st.status("Assessing the revised dissertation…", expanded=True) as status:
                    st.session_state["revised_assessment"] = generate_revised_dissertation_assessment(
                        st.session_state["revised_dissertation_text"],
                        dissertation_title,
                        criteria_bundle,
                        prior_final_review,
                        chair_concerns,
                    )
                    status.update(label="Revised assessment ready.", state="complete")

    with col_b2:
        if st.button("Evaluate Candidate Response", width="stretch"):
            if "candidate_response_text" not in st.session_state:
                st.error("Upload and process the candidate response letter first.")
            elif "revised_dissertation_text" not in st.session_state:
                st.error("Upload and process the revised dissertation first.")
            elif not criteria_bundle_parts:
                st.error("Generate per-chapter criteria in the chapter tabs first.")
            else:
                with st.status("Evaluating candidate response…", expanded=True) as status:
                    st.session_state["candidate_response_evaluation"] = evaluate_candidate_response(
                        st.session_state["candidate_response_text"],
                        prior_final_review,
                        st.session_state["revised_dissertation_text"],
                        criteria_bundle,
                        chair_concerns,
                    )
                    status.update(label="Response evaluation ready.", state="complete")

    if "revised_assessment" in st.session_state:
        with st.expander("🧭 Revised Dissertation Assessment", expanded=True):
            st.markdown(st.session_state["revised_assessment"])
        st.download_button(
            "Download Revised Assessment",
            data=st.session_state["revised_assessment"],
            file_name=f"{(dissertation_title or 'dissertation').replace(' ', '_')}_revised_assessment.txt",
            mime="text/plain",
            key="dl_revised_assessment",
        )

    if "candidate_response_evaluation" in st.session_state:
        with st.expander("📊 Candidate Response Adequacy", expanded=True):
            st.markdown(st.session_state["candidate_response_evaluation"])
        st.download_button(
            "Download Response Evaluation",
            data=st.session_state["candidate_response_evaluation"],
            file_name=f"{(dissertation_title or 'dissertation').replace(' ', '_')}_response_evaluation.txt",
            mime="text/plain",
            key="dl_response_evaluation",
        )


with analytics_tab:
    st.header("📈 Examination Analytics")

    col1, col2, col3 = st.columns(3)
    with col1:
        st.metric("Dissertation", dissertation_title if dissertation_title else "—")
    with col2:
        st.metric("Examination Focus", chair_concerns.strip() if chair_concerns.strip() else "Not specified")
    with col3:
        st.metric("Last Review Generated", st.session_state.get("review_date", "—"))

    if "extracted_texts" in st.session_state and st.session_state.extracted_texts:
        total_words = sum(len(t.split()) for t in st.session_state.extracted_texts.values())
        col_w1, col_w2 = st.columns(2)
        with col_w1:
            st.metric("Files processed", len(st.session_state.extracted_texts))
        with col_w2:
            st.metric("Total word count", f"{total_words:,}")

    st.markdown("---")
    st.subheader("Duplication Ledger")
    st.caption(
        "Every issue raised, and the chapters that referenced it again instead of restating it."
    )

    recurrences = _recurrence_map()
    ledger_rows = []
    for n in range(1, 7):
        for row in _parse_ledger(st.session_state.get(f"ch{n}_ledger", "")):
            seen_again = recurrences.get(row["ID"], [])
            ledger_rows.append(
                {
                    **row,
                    "Raised in": n,
                    "Recurs in": ", ".join(str(c) for c in seen_again) if seen_again else "—",
                }
            )

    if ledger_rows:
        st.dataframe(
            ledger_rows,
            width="stretch",
            hide_index=True,
            column_order=("ID", "Raised in", "Recurs in", "Severity", "Theme", "Issue"),
        )
        collapsed = sum(len(recurrences.get(r["ID"], [])) for r in ledger_rows)
        st.caption(
            f"{len(ledger_rows)} distinct issue{'s' if len(ledger_rows) != 1 else ''} · "
            f"{collapsed} restatement{'s' if collapsed != 1 else ''} avoided across chapters."
        )
    else:
        st.caption("No chapter reviews generated yet.")

    st.markdown("---")
    st.subheader("Artifact Checklist")

    def _row(label: str, present: bool):
        st.write(("✅ " if present else "⬜️ ") + label)

    for n in range(1, 7):
        name, icon, _ = CHAPTER_DESCRIPTIONS[n]
        _row(
            f"{icon} Chapter {n} ({name}) — Criteria",
            f"ch{n}_criteria" in st.session_state,
        )
        _row(
            f"{icon} Chapter {n} ({name}) — Review",
            f"ch{n}_review" in st.session_state,
        )
        _row(
            f"{icon} Chapter {n} ({name}) — Oral Questions",
            f"ch{n}_oral_questions" in st.session_state,
        )

    _row("📑 Dissertation Summary", "dissertation_summary" in st.session_state)
    _row("🧭 Final Comprehensive Review", "final_review" in st.session_state)
    _row("✉️ Decision Letter", "final_decision_letter" in st.session_state)
    _row("🗣️ Overarching Oral Questions", "overall_oral_questions" in st.session_state)
    _row("🔁 Revised Dissertation Assessment", "revised_assessment" in st.session_state)
    _row("📊 Candidate Response Evaluation", "candidate_response_evaluation" in st.session_state)


st.markdown("---")
st.caption(
    f"Doctoral Dissertation Examination Assistant  ·  OpenAI {OPENAI_MODEL} ({REASONING_EFFORT} reasoning)"
)
st.caption("🎓 AI-assisted examination. Final decisions must always involve human expert judgement.")
